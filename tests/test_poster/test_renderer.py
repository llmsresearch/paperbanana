"""Renderer tests: structure-golden pptx checks, no LibreOffice needed."""

from __future__ import annotations

from pathlib import Path

import pytest

from paperbanana.poster.renderer import (
    PageSizeError,
    TextOverflowError,
    check_overflow,
    measure_text_height_mm,
    render_pptx,
    required_print_scale,
    resolve_figure_placement,
    resolve_font_path,
)
from paperbanana.poster.types import (
    BBox,
    PhysicalSize,
    PosterIR,
    QRElement,
    TextElement,
    mm_to_emu,
)

from .conftest import make_panel


def test_resolve_font_path_returns_file():
    path = resolve_font_path("Helvetica")
    assert path.is_file()


def test_measure_text_height_scales_with_width():
    text = "word " * 60
    narrow = measure_text_height_mm(text, "Helvetica", 28, width_mm=100)
    wide = measure_text_height_mm(text, "Helvetica", 28, width_mm=400)
    assert narrow > wide * 2


def test_measure_text_height_counts_newlines():
    one = measure_text_height_mm("a", "Helvetica", 28, width_mm=300)
    three = measure_text_height_mm("a\nb\nc", "Helvetica", 28, width_mm=300)
    assert three == pytest.approx(one * 3)


def test_figure_placement_respects_max_height(poster_ir: PosterIR):
    panel = next(p for p in poster_ir.panels if p.id == "method")
    element = panel.elements[-1]
    placement = resolve_figure_placement(panel, element, 2000, 4000)  # tall figure
    assert panel.bbox is not None
    inner_h = panel.bbox.h_mm - 2 * 10.0
    assert placement.height_mm <= inner_h * element.max_height_frac + 0.01


def test_render_pptx_structure(poster_ir: PosterIR, tmp_path: Path):
    out = render_pptx(poster_ir, tmp_path / "poster.pptx", tmp_path / "work")
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(out))
    assert prs.slide_width == Emu(mm_to_emu(poster_ir.size.width_mm))
    assert prs.slide_height == Emu(mm_to_emu(poster_ir.size.height_mm))
    slide = prs.slides[0]
    texts = [
        shape.text_frame.text
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text
    ]
    joined = "\n".join(texts)
    assert "A Sample Paper Title" in joined
    assert "Ada Lovelace" in joined
    assert any("Step one" in t for t in texts)
    # 1 background + 4 panel rects + textboxes + 1 picture
    pictures = [s for s in slide.shapes if s.shape_type == 13]  # PICTURE
    assert len(pictures) == 1


def test_render_pptx_renders_qr(poster_ir: PosterIR, tmp_path: Path):
    data = poster_ir.model_dump()
    data["panels"][3]["elements"].append(
        QRElement(url="https://arxiv.org/abs/2601.23265", label="Paper").model_dump()
    )
    ir = PosterIR(**data)
    out = render_pptx(ir, tmp_path / "poster.pptx", tmp_path / "work")
    from pptx import Presentation

    slide = Presentation(str(out)).slides[0]
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    assert len(pictures) == 2  # figure + QR


def test_overflow_raises(poster_ir: PosterIR):
    data = poster_ir.model_dump()
    data["panels"][2]["elements"] = [
        TextElement(level="body", content="lorem ipsum " * 600).model_dump()
    ]
    ir = PosterIR(**data)
    with pytest.raises(TextOverflowError, match="results"):
        check_overflow(ir)


def _cvpr_ir(style_tokens, print_scale: int) -> PosterIR:
    return PosterIR(
        venue="cvpr",
        venue_spec_year=2025,
        size=PhysicalSize(width_mm=2133.6, height_mm=1066.8),  # 84in > 56in pptx cap
        orientation="landscape",
        print_scale=print_scale,
        style=style_tokens,
        paper_title="X",
        panels=[
            make_panel(
                "header",
                0,
                role="header",
                bbox=BBox(x_mm=10, y_mm=10, w_mm=2000, h_mm=150),
                elements=[TextElement(level="title", content="X")],
            ),
            make_panel("body", 1, bbox=BBox(x_mm=10, y_mm=180, w_mm=500, h_mm=800)),
        ],
    )


def test_oversized_page_raises_with_print_scale_hint(tmp_path: Path, style_tokens):
    with pytest.raises(PageSizeError, match="print_scale=2"):
        render_pptx(_cvpr_ir(style_tokens, 1), tmp_path / "poster.pptx", tmp_path / "work")


def test_print_scale_renders_half_size_page(tmp_path: Path, style_tokens):
    assert required_print_scale(2133.6, 1066.8) == 2
    out = render_pptx(_cvpr_ir(style_tokens, 2), tmp_path / "poster.pptx", tmp_path / "work")
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(out))
    assert prs.slide_width == Emu(mm_to_emu(2133.6 / 2))
    assert prs.slide_height == Emu(mm_to_emu(1066.8 / 2))
    # Title font is halved on the design page (printed at 200%).
    slide = prs.slides[0]
    title_runs = [
        run
        for shape in slide.shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
        if run.text == "X"
    ]
    assert title_runs and title_runs[0].font.size.pt == pytest.approx(72 / 2)


def test_unplaced_ir_raises(poster_ir: PosterIR, tmp_path: Path):
    data = poster_ir.model_dump()
    for p in data["panels"]:
        p["bbox"] = None
    ir = PosterIR(**data)
    with pytest.raises(ValueError, match="layout solver"):
        render_pptx(ir, tmp_path / "poster.pptx", tmp_path / "work")
